from __future__ import annotations

import io

import fitz
from playwright.async_api import Browser, Error as PlaywrightError, TimeoutError as PlaywrightTimeoutError
from pptx import Presentation
from pptx.util import Inches


class HTMLConversionError(Exception):
    """Raised when Chromium cannot render the supplied HTML."""


RENDER_TIMEOUT_MS = 60_000
PPTX_RENDER_DPI = 144


async def render_html_to_pdf(browser: Browser, html: str) -> bytes:
    context = await browser.new_context()
    page = await context.new_page()
    page.set_default_timeout(RENDER_TIMEOUT_MS)
    try:
        await page.set_content(html, wait_until="load", timeout=RENDER_TIMEOUT_MS)
        try:
            await page.wait_for_load_state("networkidle", timeout=15_000)
        except PlaywrightTimeoutError:
            # Some slide templates keep a connection open. Loaded fonts and
            # images below are the assets that matter to the rendered file.
            pass
        await page.evaluate(
            """async () => {
                if (document.fonts?.ready) await document.fonts.ready;
                await Promise.all(Array.from(document.images).map((image) => {
                    if (image.complete) return Promise.resolve();
                    return new Promise((resolve) => {
                        image.addEventListener('load', resolve, { once: true });
                        image.addEventListener('error', resolve, { once: true });
                    });
                }));
            }"""
        )
        await page.emulate_media(media="print")
        return await page.pdf(print_background=True, prefer_css_page_size=True)
    except (PlaywrightError, PlaywrightTimeoutError) as exc:
        raise HTMLConversionError(f"Chromium could not render the HTML: {exc}") from exc
    finally:
        await context.close()


def pdf_to_pptx(pdf: bytes) -> bytes:
    """Create a visually faithful PPTX with one rendered PDF page per slide."""
    try:
        document = fitz.open(stream=pdf, filetype="pdf")
    except Exception as exc:
        raise HTMLConversionError("The rendered PDF could not be opened") from exc

    try:
        if document.page_count == 0:
            raise HTMLConversionError("The HTML rendered without any pages")

        first_page = document[0]
        width_points, height_points = first_page.rect.width, first_page.rect.height
        presentation = Presentation()
        presentation.slide_width = Inches(width_points / 72)
        presentation.slide_height = Inches(height_points / 72)
        blank_layout = presentation.slide_layouts[6]

        scale = PPTX_RENDER_DPI / 72
        matrix = fitz.Matrix(scale, scale)
        for page in document:
            pixmap = page.get_pixmap(matrix=matrix, alpha=False)
            image = io.BytesIO(pixmap.tobytes("png"))
            slide = presentation.slides.add_slide(blank_layout)
            slide.shapes.add_picture(
                image,
                0,
                0,
                width=presentation.slide_width,
                height=presentation.slide_height,
            )

        # python-pptx starts with no slides, so no template slide needs removal.
        output = io.BytesIO()
        presentation.save(output)
        return output.getvalue()
    finally:
        document.close()
